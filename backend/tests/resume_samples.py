"""Builds synthetic resume files in memory for the Resume Parsing tests.

All people, emails and organisations here are made up.
"""

import io
from typing import List

import docx
from pypdf import PdfWriter

FAKE_RESUME_LINES: List[str] = [
    "Jordan Lee",
    "jordan.lee@example.com | +1 555 010 0000",
    "Backend-focused graduate who enjoys building reliable services",
    "Education",
    "B.Tech in Computer Science, Example Institute of Technology, 2021 - 2025, CGPA 8.4 / 10",
    "Technical Skills",
    "Python, Java, SQL (basic), Docker, Git, Linux basics, Figma",
    "Projects",
    "- Inventory API - built a REST API in Python with a PostgreSQL database and unit testing",
    "- Built a chat server in Java using multithreading for concurrent clients",
    "Experience",
    "- Software intern at Example Corp - added CI/CD checks to the deployment pipeline",
    "Hackathons",
    "- Winner, Example City Hackathon - built a delivery route planner",
    "Certifications",
    "AWS Cloud Practitioner",
    "Achievements",
    "Mentored first-year students in a weekly coding club",
    "Preferred Roles",
    "Backend Engineer, Software Engineer",
]


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines: List[str]) -> bytes:
    """A minimal one-page, text-based PDF (Helvetica, one line per entry)."""
    ops = ["BT", "/F1 11 Tf", "14 TL", "50 750 Td"]
    for i, line in enumerate(lines):
        ops.append(("" if i == 0 else "T* ") + f"({_pdf_escape(line)}) Tj")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1") if lines else b""

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()


def make_encrypted_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt(user_password="secret", owner_password="owner")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(lines: List[str]) -> bytes:
    """A DOCX where "- " lines become real Word bullet paragraphs."""
    document = docx.Document()
    for line in lines:
        if line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
