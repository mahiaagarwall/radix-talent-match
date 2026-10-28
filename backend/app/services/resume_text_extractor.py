"""Resume file validation and text extraction (Owner: Mahi Aggarwal).

Turns an uploaded PDF or DOCX (as bytes) into a list of text lines. Nothing is
written to disk and no resume text is logged.

Every problem with the file raises ResumeFileError, which carries:
- code         short machine-readable reason, also sent to the frontend
- message      a sentence that is safe to show to the user
- status_code  the HTTP status the API should return
"""

import io
import os
from typing import List

# Same limit as the frontend (frontend/src/app/config.js → MAX_UPLOAD_BYTES).
DEFAULT_MAX_UPLOAD_MB = 10

ALLOWED_EXTENSIONS = {"pdf", "docx"}

# Browsers are not consistent about MIME types (Windows often sends
# application/octet-stream), so these are accepted alongside the real types.
# The file signature check below is what actually proves the format.
ALLOWED_CONTENT_TYPES = {
    "pdf": {"application/pdf", "application/x-pdf", "application/octet-stream", ""},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
        "application/octet-stream",
        "",
    },
}

# Paragraph styles python-docx reports for bulleted/numbered list items.
DOCX_LIST_STYLE_WORDS = ("list",)


class ResumeFileError(Exception):
    def __init__(self, code: str, message: str, status_code: int = 422):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


def max_upload_bytes() -> int:
    """Upload limit in bytes. Override with MAX_UPLOAD_MB in backend/.env."""
    try:
        mb = float(os.getenv("MAX_UPLOAD_MB", DEFAULT_MAX_UPLOAD_MB))
    except ValueError:
        mb = DEFAULT_MAX_UPLOAD_MB
    return int(mb * 1024 * 1024)


def file_extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    return name.rsplit(".", 1)[-1] if "." in name else ""


def validate_upload(filename: str, content_type: str, data: bytes) -> str:
    """Checks extension, MIME type, size and file signature. Returns "pdf" or "docx"."""
    if not filename:
        raise ResumeFileError("missing", "No file was uploaded. Choose a PDF or DOCX resume.", 400)

    kind = file_extension(filename)
    if kind == "doc":
        raise ResumeFileError(
            "unsupported",
            "Legacy .doc files are not supported. Save the document as .docx or PDF and try again.",
            415,
        )
    if kind not in ALLOWED_EXTENSIONS:
        raise ResumeFileError("unsupported", f'"{filename}" is not a supported file. Upload a PDF or DOCX document.', 415)

    mime = (content_type or "").split(";")[0].strip().lower()
    if mime not in ALLOWED_CONTENT_TYPES[kind]:
        raise ResumeFileError("unsupported", f'"{filename}" was sent as "{mime}", which does not match a {kind.upper()} file.', 415)

    if not data:
        raise ResumeFileError("unreadable", f'"{filename}" is empty and cannot be read.', 422)

    limit = max_upload_bytes()
    if len(data) > limit:
        raise ResumeFileError("too_large", f'"{filename}" is larger than the {limit // (1024 * 1024)} MB limit.', 413)

    # Signature ("magic bytes") check: a renamed or corrupted file is rejected here.
    if kind == "pdf" and not data[:1024].lstrip().startswith(b"%PDF"):
        raise ResumeFileError("unreadable", f'"{filename}" has a .pdf extension but is not a valid PDF file.', 422)
    if kind == "docx" and not data.startswith(b"PK"):
        raise ResumeFileError("unreadable", f'"{filename}" has a .docx extension but is not a valid Word document.', 422)

    return kind


def _extract_pdf_lines(data: bytes) -> List[str]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # Some PDFs are "encrypted" with an empty password just to block editing.
            try:
                opened = reader.decrypt("")
            except Exception:
                opened = 0
            if not opened:
                raise ResumeFileError(
                    "encrypted",
                    "This PDF is password-protected. Remove the password (or export a new PDF) and upload it again.",
                    422,
                )
        pages = [page.extract_text() or "" for page in reader.pages]
    except ResumeFileError:
        raise
    except Exception as exc:  # a damaged PDF can make pypdf raise almost any error type
        raise ResumeFileError("unreadable", "This PDF is damaged or could not be read.", 422) from exc

    lines = [line for page in pages for line in page.splitlines()]
    if not any(line.strip() for line in lines):
        raise ResumeFileError(
            "ocr_required",
            "No selectable text was found in this PDF. It is probably a scanned image; "
            "OCR is not supported yet, so upload a text-based PDF or a DOCX file.",
            422,
        )
    return lines


def _extract_docx_lines(data: bytes) -> List[str]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # bad ZIP, missing document.xml, invalid XML, ...
        raise ResumeFileError("unreadable", "This Word document is damaged or could not be read.", 422) from exc

    lines: List[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "").lower()
        # Mark list items with a bullet so the parser treats PDF and DOCX the same way.
        if any(word in style for word in DOCX_LIST_STYLE_WORDS):
            text = f"• {text}"
        lines.extend(text.splitlines())

    # Some resume templates put everything inside tables.
    for table in document.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in seen:  # merged cells repeat their text
                    seen.add(text)
                    lines.extend(text.splitlines())

    if not any(line.strip() for line in lines):
        raise ResumeFileError("empty", "This Word document does not contain any text.", 422)
    return lines


def extract_text_lines(data: bytes, kind: str) -> List[str]:
    """Returns the document's text as a list of lines (blank lines removed)."""
    if kind == "pdf":
        lines = _extract_pdf_lines(data)
    elif kind == "docx":
        lines = _extract_docx_lines(data)
    else:
        raise ResumeFileError("unsupported", "Upload a PDF or DOCX document.", 415)
    return [line.strip() for line in lines if line.strip()]
