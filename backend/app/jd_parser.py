"""
app/jd_parser.py — PDF and DOCX text extraction for JD Analyzer.

Responsibilities:
  - Accept raw file bytes and a MIME type / filename.
  - Extract meaningful text from PDF and DOCX documents.
  - Reject unsupported formats, empty content, and corrupted files.
  - Enforce upload-size limits.
  - Never produce empty text silently.
"""

from __future__ import annotations

import io
import logging
from typing import Tuple

from app.config import get_settings

logger = logging.getLogger(__name__)

SUPPORTED_MIME_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    # Some browsers send this for .docx
    "application/octet-stream": None,  # Resolve by extension
}

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class ParserError(Exception):
    """Raised for all recoverable parsing failures."""

    def __init__(self, message: str, user_message: str | None = None):
        super().__init__(message)
        self.user_message = user_message or message


def _resolve_format(filename: str, content_type: str) -> str:
    """
    Return "pdf" or "docx" by checking the file extension first, then MIME type.
    Raises ParserError for unsupported types.
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return "pdf"
    if lower.endswith(".docx"):
        return "docx"

    fmt = SUPPORTED_MIME_TYPES.get(content_type)
    if fmt:
        return fmt

    raise ParserError(
        f"Unsupported file type: extension='{filename}', content_type='{content_type}'",
        user_message=(
            f"Unsupported file type. Please upload a PDF (.pdf) or "
            f"Word document (.docx). Received: '{filename}'."
        ),
    )


def _extract_pdf_text(data: bytes) -> str:
    """Extract text from PDF bytes using pypdf."""
    try:
        import pypdf  # lazy import — only needed for PDF paths
    except ImportError:
        raise ParserError(
            "pypdf is not installed.",
            user_message="PDF parsing is unavailable. Install the 'pypdf' package.",
        )

    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ParserError(
            f"PDF could not be opened: {exc}",
            user_message="The PDF file appears to be corrupted or password-protected.",
        )

    if len(reader.pages) == 0:
        raise ParserError(
            "PDF has zero pages.",
            user_message="The uploaded PDF contains no pages.",
        )

    parts: list[str] = []
    for page_num, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
            parts.append(text)
        except Exception as exc:
            logger.warning("Page %d extraction failed: %s", page_num, exc)

    combined = "\n".join(parts).strip()

    if not combined:
        raise ParserError(
            "PDF yielded empty text (possibly a scanned/image-only PDF).",
            user_message=(
                "No text could be extracted from this PDF. "
                "It may be a scanned document that requires OCR, "
                "which is not currently supported. "
                "Please provide a text-based PDF or a DOCX file."
            ),
        )

    return combined


def _extract_docx_text(data: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        import docx  # lazy import
    except ImportError:
        raise ParserError(
            "python-docx is not installed.",
            user_message="DOCX parsing is unavailable. Install the 'python-docx' package.",
        )

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParserError(
            f"DOCX could not be opened: {exc}",
            user_message="The DOCX file appears to be corrupted or is not a valid Word document.",
        )

    paragraphs: list[str] = []

    # Main body paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Tables (each cell's text on its own line)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in paragraphs:
                    paragraphs.append(cell_text)

    combined = "\n".join(paragraphs).strip()

    if not combined:
        raise ParserError(
            "DOCX yielded empty text.",
            user_message="No text could be extracted from the DOCX file. The document may be empty.",
        )

    return combined


def extract_text(data: bytes, filename: str, content_type: str = "") -> Tuple[str, str]:
    """
    Extract plain text from a PDF or DOCX file.

    Args:
        data          : Raw file bytes.
        filename      : Original filename (used for extension-based detection).
        content_type  : MIME type reported by the client.

    Returns:
        A (text, format) tuple where format is "pdf" or "docx".

    Raises:
        ParserError: For all recoverable extraction failures.
    """
    settings = get_settings()

    # ── Size guard ───────────────────────────────────────────────────────────
    if len(data) == 0:
        raise ParserError(
            "Uploaded file is empty (0 bytes).",
            user_message="The uploaded file is empty. Please upload a non-empty PDF or DOCX.",
        )

    if len(data) > settings.max_upload_bytes:
        limit_mb = settings.max_upload_bytes / (1024 * 1024)
        raise ParserError(
            f"File size {len(data)} bytes exceeds limit {settings.max_upload_bytes} bytes.",
            user_message=f"File too large. Maximum allowed size is {limit_mb:.0f} MB.",
        )

    # ── Format detection ──────────────────────────────────────────────────────
    fmt = _resolve_format(filename, content_type)

    # ── Extraction ────────────────────────────────────────────────────────────
    if fmt == "pdf":
        text = _extract_pdf_text(data)
    else:
        text = _extract_docx_text(data)

    logger.info("Extracted %d characters from %s (%s)", len(text), filename, fmt)
    return text, fmt
